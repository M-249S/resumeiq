from io import BytesIO

from docx import Document
from docx.shared import Pt


def generate_cover_letter_docx(
    cover_letter: str,
) -> BytesIO:
    """
    Generate a DOCX cover letter from Markdown/plain text.
    """

    document = Document()

    style = document.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    for line in cover_letter.split("\n"):

        line = line.strip()

        if not line:
            document.add_paragraph()
            continue

        # Markdown Heading
        if line.startswith("# "):
            document.add_heading(
                line.replace("# ", ""),
                level=1,
            )

        elif line.startswith("## "):
            document.add_heading(
                line.replace("## ", ""),
                level=2,
            )

        elif line.startswith("### "):
            document.add_heading(
                line.replace("### ", ""),
                level=3,
            )

        else:
            document.add_paragraph(line)

    buffer = BytesIO()

    document.save(buffer)

    buffer.seek(0)

    return buffer
